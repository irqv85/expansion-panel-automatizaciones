#!/usr/bin/env python3
# GB Freshworks CX Cadence Deck — self-contained generator.
# Usage: python build_deck.py <deals.xlsx> <assets.xlsx> <vtiger.json> <out.html> [YYYY-MM-DD]
# vtiger.json = {"POT": {"<Deal ID>": {"cf948":..,"plan":..,"agents":..,"id":"5x.."}}, "Q": {"5x..": [[quote_no,total,stage],...]}}
import json, html, re, sys, datetime, os
import openpyxl
F1, F2, VT, OUT = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]
TODAY = sys.argv[5] if len(sys.argv)>5 else datetime.date.today().isoformat()
CUR_MONTH = TODAY[:7]
_dt = datetime.datetime.strptime(TODAY,"%Y-%m-%d")
DATE = _dt.strftime("%B ")+str(_dt.day)+_dt.strftime(", %Y")
TEMPLATE = open(os.path.join(os.path.dirname(os.path.abspath(__file__)),"template.html"),encoding="utf-8").read()
LOGO_W = '<svg width="408" height="86" viewBox="0 0 408 86" fill="none" xmlns="http://www.w3.org/2000/svg"> <path d="M408 54.1533C408 60.9129 402.302 66 393.027 66C383.155 66 376.53 60.216 376 52.6899H386.667C386.932 55.4077 389.781 57.2195 392.894 57.2195C395.81 57.2195 397.029 55.2977 397.029 53.5556C397.029 47.2838 376 52.2857 376 38C376 31.3798 382.559 26 392.232 26C401.772 26 407.271 31.6678 408 39.3333H397.333C397.002 36.6852 395.081 34.8502 391.901 34.8502C389.251 34.8502 386.971 36.5629 386.971 38.4444C386.971 44.6465 407.801 39.6585 408 54.1533Z" fill="#FFFFFF"></path> <path fill-rule="evenodd" clip-rule="evenodd" d="M352 66V46V45.9717V26H362V28.6802C364.941 26.9755 368.357 26 372 26V36C366.473 36 362 40.4665 362 46V66H352Z" fill="#FFFFFF"></path> <path fill-rule="evenodd" clip-rule="evenodd" d="M348 46C348 57.0457 339.046 66 328 66C316.954 66 308 57.0457 308 46C308 34.9543 316.954 26 328 26C339.046 26 348 34.9543 348 46ZM328 56C333.523 56 338 51.5228 338 46C338 40.4772 333.523 36 328 36C322.477 36 318 40.4772 318 46C318 51.5228 322.477 56 328 56Z" fill="#FFFFFF"></path> <path d="M304 54.1533C304 60.9129 298.302 66 289.027 66C279.155 66 272.53 60.216 272 52.6899H282.667C282.932 55.4077 285.781 57.2195 288.894 57.2195C291.81 57.2195 293.029 55.2977 293.029 53.5556C293.029 47.2838 272 52.2857 272 38C272 31.3798 278.559 26 288.232 26C297.772 26 303.271 31.6678 304 39.3333H293.333C293.002 36.6852 291.081 34.8502 287.901 34.8502C285.251 34.8502 282.971 36.5629 282.971 38.4444C282.971 44.6465 303.801 39.6585 304 54.1533Z" fill="#FFFFFF"></path> <circle cx="261" cy="12" r="6" fill="#FFFFFF"></circle> <path d="M266 26H256V66H266V26Z" fill="#FFFFFF"></path> <path d="M237.778 66L252 26H241.361L232 52.6841L222.645 26H212L226.222 66H237.778Z" fill="#FFFFFF"></path> <path fill-rule="evenodd" clip-rule="evenodd" d="M208 6H198V28.6791C194.913 26.7289 191.922 26 188 26C176.954 26 168 34.9543 168 46C168 57.0457 176.954 66 188 66C191.643 66 195.058 65.0261 198 63.3243V66H208V46V6ZM198 46C198 40.4772 193.523 36 188 36C182.477 36 178 40.4772 178 46C178 51.5228 182.477 56 188 56C193.523 56 198 51.5228 198 46Z" fill="#FFFFFF"></path> <path fill-rule="evenodd" clip-rule="evenodd" d="M154 28.5V26H164V66H154V63.5C151.055 65.2041 147.624 66 143.981 66C132.945 66 124 57.0464 124 46C124 34.9537 132.945 26 143.981 26C147.624 26 151.055 26.7959 154 28.5ZM144 56C149.523 56 154 51.5228 154 46C154 40.4772 149.523 36 144 36C138.477 36 134 40.4772 134 46C134 51.5228 138.477 56 144 56Z" fill="#FFFFFF"></path> <path fill-rule="evenodd" clip-rule="evenodd" d="M80 6H70V46V66H80V63.3243C82.9417 65.0261 86.3571 66 90 66C101.046 66 110 57.0457 110 46C110 34.9543 101.046 26 90 26C86.3571 26 82.9417 26.9739 80 28.6757V6ZM80 46C80 51.5228 84.4772 56 90 56C95.5228 56 100 51.5228 100 46C100 40.4772 95.5228 36 90 36C84.4772 36 80 40.4772 80 46Z" fill="#FFFFFF"></path> <path fill-rule="evenodd" clip-rule="evenodd" d="M65.9949 46.1337H65.9975V65.8967H65.9924C65.9939 65.9282 65.9954 65.9589 65.9964 65.9897C65.9971 66.0106 65.9975 66.0316 65.9975 66.053C65.9975 77.0686 57.0451 86 45.9987 86C34.9549 86 26 77.0706 26 66.053C26.0045 66.3535 26.0068 66.5068 26.007 66.5068C26.0071 66.5068 26.0048 66.3403 26 66H36C36 71.5092 40.4769 76 46 76C51.5232 76 55.9981 71.5622 55.9981 66.053C55.9981 66.0001 55.993 65.9471 55.993 65.8941H55.9981V63.3333C53.0569 65.0334 49.6442 66 46 66C34.9562 66 26 56.9648 26 45.9471C26 34.9293 34.9549 26 46.0013 26C49.6435 26 53.0585 26.9712 56 28.6682V26H66V45.9471V46C65.9996 46.0135 65.998 46.0269 65.9975 46.0403C65.9962 46.0712 65.9949 46.1022 65.9949 46.1337ZM35.9994 45.9471C35.9994 51.4559 40.4769 56 46 56C51.5232 56 56.0006 51.4559 55.9981 45.9471C55.9981 40.4381 51.4282 36 46 36C40.5718 36 35.9994 40.4381 35.9994 45.9471Z" fill="#FFFFFF"></path> <path d="M19.4999 0H20.5C22.4885 9.79565 30.2045 17.5115 40 19.4999V20.5C30.2045 22.4885 22.4885 30.2045 20.5 40H19.4999C17.5115 30.2045 9.79565 22.4885 0 20.5V19.4999C9.79565 17.5115 17.5115 9.79565 19.4999 0Z" fill="#EA018B"></path> </svg>'
LOGO_D = '<svg width="408" height="86" viewBox="0 0 408 86" fill="none" xmlns="http://www.w3.org/2000/svg"> <path d="M408 54.1533C408 60.9129 402.302 66 393.027 66C383.155 66 376.53 60.216 376 52.6899H386.667C386.932 55.4077 389.781 57.2195 392.894 57.2195C395.81 57.2195 397.029 55.2977 397.029 53.5556C397.029 47.2838 376 52.2857 376 38C376 31.3798 382.559 26 392.232 26C401.772 26 407.271 31.6678 408 39.3333H397.333C397.002 36.6852 395.081 34.8502 391.901 34.8502C389.251 34.8502 386.971 36.5629 386.971 38.4444C386.971 44.6465 407.801 39.6585 408 54.1533Z" fill="#25235A"></path> <path fill-rule="evenodd" clip-rule="evenodd" d="M352 66V46V45.9717V26H362V28.6802C364.941 26.9755 368.357 26 372 26V36C366.473 36 362 40.4665 362 46V66H352Z" fill="#25235A"></path> <path fill-rule="evenodd" clip-rule="evenodd" d="M348 46C348 57.0457 339.046 66 328 66C316.954 66 308 57.0457 308 46C308 34.9543 316.954 26 328 26C339.046 26 348 34.9543 348 46ZM328 56C333.523 56 338 51.5228 338 46C338 40.4772 333.523 36 328 36C322.477 36 318 40.4772 318 46C318 51.5228 322.477 56 328 56Z" fill="#25235A"></path> <path d="M304 54.1533C304 60.9129 298.302 66 289.027 66C279.155 66 272.53 60.216 272 52.6899H282.667C282.932 55.4077 285.781 57.2195 288.894 57.2195C291.81 57.2195 293.029 55.2977 293.029 53.5556C293.029 47.2838 272 52.2857 272 38C272 31.3798 278.559 26 288.232 26C297.772 26 303.271 31.6678 304 39.3333H293.333C293.002 36.6852 291.081 34.8502 287.901 34.8502C285.251 34.8502 282.971 36.5629 282.971 38.4444C282.971 44.6465 303.801 39.6585 304 54.1533Z" fill="#25235A"></path> <circle cx="261" cy="12" r="6" fill="#25235A"></circle> <path d="M266 26H256V66H266V26Z" fill="#25235A"></path> <path d="M237.778 66L252 26H241.361L232 52.6841L222.645 26H212L226.222 66H237.778Z" fill="#25235A"></path> <path fill-rule="evenodd" clip-rule="evenodd" d="M208 6H198V28.6791C194.913 26.7289 191.922 26 188 26C176.954 26 168 34.9543 168 46C168 57.0457 176.954 66 188 66C191.643 66 195.058 65.0261 198 63.3243V66H208V46V6ZM198 46C198 40.4772 193.523 36 188 36C182.477 36 178 40.4772 178 46C178 51.5228 182.477 56 188 56C193.523 56 198 51.5228 198 46Z" fill="#25235A"></path> <path fill-rule="evenodd" clip-rule="evenodd" d="M154 28.5V26H164V66H154V63.5C151.055 65.2041 147.624 66 143.981 66C132.945 66 124 57.0464 124 46C124 34.9537 132.945 26 143.981 26C147.624 26 151.055 26.7959 154 28.5ZM144 56C149.523 56 154 51.5228 154 46C154 40.4772 149.523 36 144 36C138.477 36 134 40.4772 134 46C134 51.5228 138.477 56 144 56Z" fill="#25235A"></path> <path fill-rule="evenodd" clip-rule="evenodd" d="M80 6H70V46V66H80V63.3243C82.9417 65.0261 86.3571 66 90 66C101.046 66 110 57.0457 110 46C110 34.9543 101.046 26 90 26C86.3571 26 82.9417 26.9739 80 28.6757V6ZM80 46C80 51.5228 84.4772 56 90 56C95.5228 56 100 51.5228 100 46C100 40.4772 95.5228 36 90 36C84.4772 36 80 40.4772 80 46Z" fill="#25235A"></path> <path fill-rule="evenodd" clip-rule="evenodd" d="M65.9949 46.1337H65.9975V65.8967H65.9924C65.9939 65.9282 65.9954 65.9589 65.9964 65.9897C65.9971 66.0106 65.9975 66.0316 65.9975 66.053C65.9975 77.0686 57.0451 86 45.9987 86C34.9549 86 26 77.0706 26 66.053C26.0045 66.3535 26.0068 66.5068 26.007 66.5068C26.0071 66.5068 26.0048 66.3403 26 66H36C36 71.5092 40.4769 76 46 76C51.5232 76 55.9981 71.5622 55.9981 66.053C55.9981 66.0001 55.993 65.9471 55.993 65.8941H55.9981V63.3333C53.0569 65.0334 49.6442 66 46 66C34.9562 66 26 56.9648 26 45.9471C26 34.9293 34.9549 26 46.0013 26C49.6435 26 53.0585 26.9712 56 28.6682V26H66V45.9471V46C65.9996 46.0135 65.998 46.0269 65.9975 46.0403C65.9962 46.0712 65.9949 46.1022 65.9949 46.1337ZM35.9994 45.9471C35.9994 51.4559 40.4769 56 46 56C51.5232 56 56.0006 51.4559 55.9981 45.9471C55.9981 40.4381 51.4282 36 46 36C40.5718 36 35.9994 40.4381 35.9994 45.9471Z" fill="#25235A"></path> <path d="M19.4999 0H20.5C22.4885 9.79565 30.2045 17.5115 40 19.4999V20.5C30.2045 22.4885 22.4885 30.2045 20.5 40H19.4999C17.5115 30.2045 9.79565 22.4885 0 20.5V19.4999C9.79565 17.5115 17.5115 9.79565 19.4999 0Z" fill="#EA018B"></path> </svg>'
vt = json.load(open(VT)); POT = vt.get("POT",{}); Q = {k:v for k,v in vt.get("Q",{}).items()}

def esc(s): return html.escape(s or "")
def short(s,n=175):
    s=(s or "").replace("\n"," ").replace("\xa0"," ").strip(); return (s[:n]+"…") if len(s)>n else s
def cell(r,ix,n):
    i=ix.get(n); x=r[i] if i is not None else None
    return "" if x is None else str(x).strip()

# ---------- parse deals (file1) ----------
wb=openpyxl.load_workbook(F1,data_only=True); ws=wb.active
rows=list(ws.iter_rows(values_only=True)); hdr=[str(c).strip() if c else "" for c in rows[0]]; ix={h:i for i,h in enumerate(hdr)}
D42=re.compile(r"device\s?42|\bd42\b|assets?\s*pack|\bassets?\b",re.I)
AI=re.compile(r"\bbots?\b|bot session|freddy|copilot",re.I)
CX=re.compile(r"\bFSAS\b|freshsales\s*suite|\bFDO\b|fd[\s\-/]*omni|freshdesk\s*omni|\bomni\b|\bFSA\b|freshsales|\bFCH\b|fchat|freshchat|\bFCA\b|fcaller|freshcaller|\bFM\b|freshmarketer|\bFD\b|freshdesk",re.I)
EXr=re.compile(r"\bFS\b|freshservice",re.I)
def cats(name,pf,typ):
    t=f"{name} {pf}"; c=set()
    if D42.search(t): c.add("D42")
    if AI.search(t): c.add("AI")
    if typ=="Payment Frequency": c.add("PF"); return sorted(c)
    if CX.search(t) and not EXr.search(name): c.add("CX")
    if EXr.search(t): c.add("EX")
    if "CX" in c and "EX" in c:
        if CX.search(name) and not EXr.search(name): c.discard("EX")
        elif EXr.search(name) and not CX.search(name): c.discard("CX")
    if not (c & {"CX","EX","D42","AI"}): c.add("Other")
    return sorted(c)
WON={"Closed Won"}; LOST={"Closed Lost"}
seen=set(); recs=[]
for r in rows[1:]:
    if all(x is None for x in r): continue
    did=cell(r,ix,"Deals Deal ID")
    if did and did in seen: continue
    if did: seen.add(did)
    nm=cell(r,ix,"Deals Deal Name"); pf=cell(r,ix,"Deals Product"); stage=cell(r,ix,"Deals Sales Stage")
    p=POT.get(did,{}); typ=p.get("cf948","New")
    seg="Expansion" if typ in ("New-Add-On","New-Cross-Sell","Payment Frequency") else "New Business"
    emp=cell(r,ix,"Organizations Employee Count"); emp="" if emp in ("None","") else emp
    try: mrr=float(cell(r,ix,"Deals Amount"))
    except: mrr=0.0
    recs.append({"did":did,"oid":cell(r,ix,"Organizations Organization ID"),"pid":p.get("id",""),
      "plan":p.get("plan",""),"agents":p.get("agents",""),"type":typ,
      "cats":cats(nm,pf,typ),"seg":seg,
      "state":"WON" if stage in WON else ("LOST" if stage in LOST else "OPEN"),
      "stage":stage,"rep":cell(r,ix,"Deals Assigned To"),"org":cell(r,ix,"Deals Organization Name") or nm,
      "deal":nm,"country":cell(r,ix,"Organizations Billing Country") or "Unknown","ecd":cell(r,ix,"Deals Expected Close Date")[:10],
      "amt":cell(r,ix,"Deals Amount"),"status":cell(r,ix,"Deals Where are we?"),"next":cell(r,ix,"Deals Next Step"),
      "pf":pf,"emp":emp,"mrr":mrr})

# ---------- installed base (file2) ----------
wb2=openpyxl.load_workbook(F2,data_only=True); ws2=wb2.active
r2=list(ws2.iter_rows(values_only=True)); h2=[str(c).strip() if c else "" for c in r2[0]]; ix2={h:i for i,h in enumerate(h2)}
FW=re.compile(r"fresh|omnichannel|customer service suite|device\s?42",re.I)
def shortp(p):
    p=re.sub(r"\s+"," ",p)
    return p.replace("Support Desk ","").replace("Freshdesk","FD").replace("FreshDesk","FD").replace("Freshservice","FS").replace("FreshService","FS").replace("Freshchat","FCH").replace("FreshSales","FSales").replace("Freshsales","FSales").replace("Omnichannel","Omni").replace("Customer Service Suite","CSS").replace(" - "," ").strip()
import collections as _c
seenA=set(); IB=_c.defaultdict(list); ibkey=set()
for r in r2[1:]:
    aid=cell(r,ix2,"Assets Asset ID")
    if aid and aid in seenA: continue
    if aid: seenA.add(aid)
    prod=cell(r,ix2,"Assets Asset Name"); status=cell(r,ix2,"Assets Customer Status"); oid=cell(r,ix2,"Organizations Organization ID")
    ld=cell(r,ix2,"Assets License Details"); mrr=cell(r,ix2,"Assets Current MRR")
    if not oid or not prod or status=="Cancelled" or not FW.search(prod): continue
    det=re.split(r"\n",ld)[0].strip() if ld else ""
    k=(oid,prod,det)
    if k in ibkey: continue
    ibkey.add(k)
    try: m=float(mrr) if mrr else 0.0
    except: m=0.0
    IB[oid].append({"p":shortp(prod),"det":det,"st":status,"mrr":round(m,2)})
IB=dict(IB)

# ---------- helpers ----------
SW={"Qualification":.10,"Demo or POC":.20,"Why Analysis":.30,"Quote or Proposal":.50,"Accepted Proposal":.65,"Negotiation":.80}
BADGE={"Negotiation":"b-neg","Quote or Proposal":"b-quo","Demo or POC":"b-demo","Why Analysis":"b-why","Accepted Proposal":"b-acc","Qualification":"b-qual","Closed Won":"b-won","PO Invoiced":"b-won"}
PORD=[("Freshsales Suite",r"\bFSAS\b|freshsales suite"),("Freshdesk Omni",r"\bFDO\b|fd[\s\-/]*omni|freshdesk omni|\bomni\b"),
("Freshsales",r"\bFSA\b|freshsales"),("Freshchat",r"\bFCH\b|fchat|freshchat"),("Freshcaller",r"\bFCA\b|fcaller|freshcaller"),
("Freshmarketer",r"\bFM\b|freshmarketer"),("Freshdesk",r"\bFD\b|freshdesk"),("Freshservice",r"\bFS\b|freshservice"),("Device42",r"device\s?42|\bd42\b"),("Freddy AI",r"freddy|copilot")]
def plabel(name,pf):
    t=f"{name} {pf}"; hits=[]
    for lab,rx in PORD:
        if re.search(rx,t,re.I) and lab not in hits: hits.append(lab)
    return " + ".join(hits[:2]) if hits else (pf or "—")
def qstage_cls(s): return "q-act" if s in ("Ready","Invoice approved") else ("q-exp" if s in ("Expired","Discarded") else "q-prog")
def money(n): return "$"+format(int(round(n)),",")
def qlist_html(pid):
    qs=Q.get(pid)
    if not qs: return '<span class="q q-no">Not quoted</span>'
    qs=sorted(qs,key=lambda x:-x[0])
    return '<div class="qlist">'+''.join(f'<span class="q {qstage_cls(st)}" title="QUO{no}">{esc((money(t) if t>0 else "draft")+" · "+st)}</span>' for no,t,st in qs)+'</div>'
CATLAB={"CX":"CX","EX":"EX","D42":"Device 42","AI":"AI","PF":"Pay freq","Other":"Add-on"}; CATCLS={"CX":"ct-cx","EX":"ct-ex","D42":"ct-d42","AI":"ct-ai","PF":"ct-pf","Other":"ct-ot"}
def chips(r):
    c=''.join(f'<span class="ctag {CATCLS[t]}">{CATLAB[t]}</span>' for t in r["cats"])
    c+=f'<span class="chip">{esc(plabel(r["deal"],r["pf"]))}</span>'
    if r["plan"]: c+=f'<span class="chip chip-p">{esc(r["plan"])}</span>'
    if r["agents"]: c+=f'<span class="chip chip-a">{esc(r["agents"])} ag</span>'
    return c
def ibblock(r):
    if r.get("seg")!="Expansion": return ""
    a=IB.get(r.get("oid"))
    if not a: return ""
    def sh(x):
        det=x["det"].split("@")[0].strip() if x.get("det") else ""
        return '<span class="ibit">'+esc(x["p"])+((" · "+esc(det)) if det else "")+'</span>'
    return '<div class="ib"><span class="iblab">Has</span>'+''.join(sh(x) for x in a[:10])+'</div>'
def is_cancclose(r):
    return r["stage"] in ("PO Invoiced","Accepted Proposal") or (r["stage"]=="Negotiation" and (r["ecd"] or "")[:7]==CUR_MONTH)

op=[r for r in recs if r["state"]=="OPEN"]; won=[r for r in recs if r["state"]=="WON"]
reps=sorted({x["rep"] for x in op}); countries=sorted({x["country"] for x in op})
def slim(r,w=False):
    o={"r":r["rep"],"o":r["org"],"c":r["cats"],"p":plabel(r["deal"],r["pf"]),"m":round(r["mrr"],2)}
    if not w: o.update({"g":"NB" if r["seg"]=="New Business" else "Exp","co":r["country"],"s":r["stage"],
      "w":round(r["mrr"]*SW.get(r["stage"],.2),2),"em":(r["ecd"] or "")[:7],"oid":r.get("oid",""),"emp":r.get("emp",""),
      "cm":(round(r["mrr"],2) if is_cancclose(r) else 0.0),"q":bool(Q.get(r["pid"]))})
    return o
DEALS=[slim(r) for r in op]; WON=[slim(r,True) for r in won]
catc=_c.Counter()
for x in op:
    for t in x["cats"]: catc[t]+=1
fwline=" · ".join(f"{CATLAB[k]} {catc[k]}" for k in ["CX","EX","D42","AI","Other"] if catc.get(k))
def rows_(rs): return sorted(rs,key=lambda x:-x["mrr"])

slides=[]
slides.append(f'<section class="slide cover">{LOGO_W}<h1>Freshworks pipeline</h1><div class="sub">CX cadence review — Iran’s team</div><div class="meta"><span>All Freshworks · MRR</span><span>CX primary · toggle EX · Device 42 · Freddy/AI</span><span>By rep · {DATE}</span></div></section>')
slides.append(f'''<section class="slide" id="execslide">{{HEAD:Executive summary|By rep — live by category}}
<div class="krow">
 <div class="kpi"><div class="k">Open pipeline</div><div class="v" id="k-open"></div><div class="d" id="k-open-d"></div></div>
 <div class="kpi"><div class="k">Can close</div><div class="v accent2" id="k-cm"></div><div class="d">PO · Accepted · Neg. this month</div></div>
 <div class="kpi kpi-won"><div class="k">Closed won</div><div class="v" id="k-won"></div><div class="d" id="k-won-d"></div></div>
 <div class="kpi"><div class="k">Quoted</div><div class="v" id="k-q"></div><div class="d">deals with a quote</div></div>
</div>
<div class="wonstrip" id="wonstrip"></div>
<div class="fwbook">Full FW open book ({len(op)} deals): {fwline}. <span class="mut">Toggle categories in the menu — every table below updates live.</span></div>
<div class="body2b"><div class="panel"><div class="ph">Pipeline by rep <span class="mut">(active categories)</span></div><div id="exec-byrep"></div></div>
 <div class="panel"><div class="ph">Talking points <span class="mut">(editable)</span></div>
  <div class="hl" contenteditable="true" data-editkey="tp1">By rep — new business first, then expansion. CX shown by default; toggle EX / Device 42 / Freddy-AI as needed.</div>
  <div class="hl" contenteditable="true" data-editkey="tp2" style="margin-top:9px">Numbers and tables update with the category toggles and country filter.</div>
  <div class="hl" contenteditable="true" data-editkey="tp3" style="margin-top:9px">Freddy/AI and bots attach to both CX and EX; Device 42 is a Freshservice add-on.</div>
 </div></div>{{FOOT}}</section>''')
def section(seg):
    big="NEW BUSINESS" if seg=="New Business" else "EXPANSION"; part="Part 1 · New business" if seg=="New Business" else "Part 2 · Expansion"
    return f'''<section class="slide secslide" data-seg="{'NB' if seg=='New Business' else 'Exp'}"><div class="bigsec">{big}</div>{{HEAD:{part}|Session overview — live}}
<div class="krow k3">
 <div class="kpi"><div class="k">Pipeline (MRR)</div><div class="v s-mrr"></div><div class="d s-mrr-d"></div></div>
 <div class="kpi"><div class="k">Can close</div><div class="v s-cm"></div><div class="d">PO · Accepted · Neg. this month</div></div>
 <div class="kpi"><div class="k">Quoted</div><div class="v s-q"></div><div class="d">deals quoted</div></div>
</div>
<div class="body2b" style="margin-top:14px">
 <div class="panel"><div class="ph">Top deals</div><div class="s-top"></div></div>
 <div class="panel"><div class="ph s-acct-h">Top accounts</div><div class="s-acct"></div><div class="ph" style="margin-top:12px">By rep</div><div class="s-byrep"></div></div>
</div>{{FOOT}}</section>'''
def rep_slide(rep,seg):
    rr=rows_([r for r in op if r["rep"]==rep and r["seg"]==seg])
    def line(r):
        b=BADGE.get(r["stage"],"b-qual"); amt=money(r["mrr"]) if r["mrr"]>0 else '<span class="pos">no amt</span>'
        od=' pos' if (r["ecd"] and r["ecd"]<TODAY) else ''
        return (f'<tr class="deal-row" data-cats="{" ".join(r["cats"])}" data-country="{esc(r["country"])}" data-mrr="{r["mrr"]}" data-w="{round(r["mrr"]*SW.get(r["stage"],.2),2)}" data-stage="{esc(r["stage"])}" data-ecdm="{(r["ecd"] or "")[:7]}">'
          f'<td><b>{esc(r["org"])}</b><div class="sub2">{esc(r["country"])}</div><div class="chips">{chips(r)}</div>{ibblock(r)}</td>'
          f'<td class="dn">{esc(r["deal"])}</td>'
          f'<td><span class="badge {b}">{esc(r["stage"])}</span><div class="sar">{amt}</div><div class="sub2{od}">{r["ecd"] or "—"}</div></td>'
          f'<td>{qlist_html(r["pid"])}</td><td class="stat" title="{esc(r["status"])}">{esc(short(r["status"]))}</td>'
          f'<td class="ns">{esc(short(r["next"],105))}</td><td contenteditable="true" class="edit" data-editkey="blk:{r["did"]}"></td></tr>')
    body='<div class="tbl-scroll"><table class="tbl deals"><thead><tr><th>Account</th><th>Deal name</th><th>Stage · MRR · ECD</th><th>Quotes</th><th>Status (CRM)</th><th>Next step</th><th>Blockers / competition</th></tr></thead><tbody>'+''.join(line(r) for r in rr)+'</tbody></table></div>'
    part="Part 1 · New business" if seg=="New Business" else "Part 2 · Expansion"
    return (f'<section class="slide repslide" data-rep="{esc(rep)}" data-seg="{seg}">{{HEAD:{part}|{esc(rep)}}}'
      f'<div class="repbar"><span class="s-cnt">—</span><span class="dot">·</span><span>Pipeline <b class="s-arr">—</b> MRR</span>'
      f'<span class="dot">·</span><span>Can close <b class="s-cm2">—</b></span><span class="dot">·</span><span>Quoted <b class="s-qn">—</b></span></div>{body}{{FOOT}}</section>')
slides.append(section("New Business"))
for rep in sorted(reps,key=lambda rp:-sum(r["mrr"] for r in op if r["rep"]==rp and r["seg"]=="New Business" and "CX" in r["cats"])):
    if any(r["rep"]==rep and r["seg"]=="New Business" for r in op): slides.append(rep_slide(rep,"New Business"))
slides.append(section("Expansion"))
for rep in sorted(reps,key=lambda rp:-sum(r["mrr"] for r in op if r["rep"]==rp and r["seg"]=="Expansion" and "CX" in r["cats"])):
    if any(r["rep"]==rep and r["seg"]=="Expansion" for r in op): slides.append(rep_slide(rep,"Expansion"))
slides.append(f'''<section class="slide summaryslide">{{HEAD:Session summary|Blockers & what we agreed}}
<div class="body2b">
 <div class="panel"><div class="ph">Blockers & competition <span class="mut">(from the deal rows)</span></div><div class="tbl-scroll" id="blkSummary"><div class="mut" style="padding:8px">Type blockers on the rep slides — they appear here.</div></div></div>
 <div class="panel"><div class="ph">Agreed in the session <span class="mut">(editable)</span></div>
  <div class="hl" contenteditable="true" data-editkey="agreements" style="min-height:120px">&nbsp;</div>
  <div class="editnote" style="font-size:11px;color:#868E96;margin-top:6px">Decisions / next steps agreed live. Auto-saved.</div>
 </div>
</div>{{FOOT}}</section>''')
out=[]; total=len(slides)
for i,sd in enumerate(slides):
    def hd(m):
        ey,t=m.group(1).split("|"); return f'<div class="eyebrow">{ey}</div><h1 class="title">{t}</h1><div class="pill"></div>'
    sd=re.sub(r"\{HEAD:([^}]*)\}",hd,sd)
    sd=sd.replace("{FOOT}",f'<div class="foot"><span>GB Advisors · confidential — partner cadence</span><span class="fl">{LOGO_D}</span><span class="pg">{i+1} / {total}</span></div>')
    out.append(sd)
SLIDES="\n".join(out)
IBMRR={oid:round(sum(x.get("mrr",0) for x in items),2) for oid,items in IB.items()}
deck={"DEALS":DEALS,"WON":WON,"COUNTRIES":countries,"IBMRR":IBMRR}
# Notas de la cadencia anterior. Viven en notas-cadencia.json, junto a este
# script, y se inyectan en el deck para que cada deal aparezca con el blocker
# que se escribio el mes pasado. Antes del 16-sep-2026 las notas solo estaban
# en el localStorage del navegador y con una clave que incluia la fecha del
# deck, asi que cada regeneracion arrancaba en blanco.
NOTAS_FILE=os.path.join(os.path.dirname(os.path.abspath(__file__)),"notas-cadencia.json")
try:
    with open(NOTAS_FILE,encoding="utf-8") as f: NOTAS=json.load(f)
    print("notas cargadas:",len(NOTAS),"de",os.path.basename(NOTAS_FILE))
except FileNotFoundError:
    NOTAS={}
except Exception as e:
    NOTAS={}; print("aviso: no se pudieron leer las notas (",e,") -- el deck sale sin ellas")
final=TEMPLATE.replace("%%SLIDES%%",SLIDES).replace("%%COUNTRIES%%",json.dumps(countries,ensure_ascii=False)).replace("%%DECKDATA%%",json.dumps(deck,ensure_ascii=False)).replace("%%DECKDATE%%",TODAY).replace("%%NOTAS%%",json.dumps(NOTAS,ensure_ascii=False))
open(OUT,"w",encoding="utf-8").write(final)
print("wrote",OUT,"| slides",total,"| open",len(op),"| won",len(won),"| IB orgs",len(IB))
