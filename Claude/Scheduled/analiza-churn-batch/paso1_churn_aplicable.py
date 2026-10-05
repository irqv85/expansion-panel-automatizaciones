#!/usr/bin/env python3
"""
Paso 1: Identificar organizaciones con churn aplicable
Replica la lógica exacta de churnValidationRows del Dashboard.html
"""

import openpyxl
from pathlib import Path
from datetime import datetime
import json
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "scripts"))
from paths import CHURN_BATCH_DIR, DOWNLOADS

def clean(v):
    """Replica clean() del Dashboard: trim y lowercase"""
    if v is None or v == '':
        return ''
    return str(v).strip().lower()

def money(v):
    """Replica money() del Dashboard: convierte a número, quitando $ y ,"""
    if v is None or v == '' or v == 0:
        return 0
    try:
        s = str(v).replace('$', '').replace(',', '')
        return float(s) or 0
    except:
        return 0

def org_key(row):
    """Replica orgKey(): usa ID si existe, si no usa Name en minúsculas"""
    org_id = clean(row.get('Organizations Organization ID', ''))
    if org_id:
        return f"id:{org_id}"
    org_name = clean(row.get('Organizations Organization Name', ''))
    if org_name:
        return f"name:{org_name}"
    return None

def load_excel_as_dicts(filepath, sheet_name=0):
    """Carga un Excel y devuelve lista de dicts (encabezados como keys)"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[sheet_name] if isinstance(sheet_name, str) else wb.worksheets[sheet_name]

    # Primera fila = encabezados
    headers = []
    for cell in ws[1]:
        headers.append(cell.value)

    rows = []
    for row_idx in range(2, ws.max_row + 1):
        row_data = {}
        for col_idx, header in enumerate(headers, start=1):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if header:
                row_data[header] = cell_value
        rows.append(row_data)

    return rows

def churn_validation_rows(churn_rows, oa_rows):
    """
    Replica churnValidationRows del Dashboard.
    Devuelve lista de dicts con info de cada fila, incluyendo si "applies"
    """
    # Armar mapa de assets activos por organización
    active_by_org = {}
    for asset in oa_rows:
        if money(asset.get('Assets Current MRR', 0)) <= 0:
            continue
        key = org_key(asset)
        if not key:
            continue
        if key not in active_by_org:
            active_by_org[key] = []
        active_by_org[key].append(asset)

    # Procesar cada fila de churn
    results = []
    for row in churn_rows:
        key = org_key(row)
        asset_name = clean(row.get('Assets Asset Name', ''))

        # Filtrar assets activos de esta org, excluyendo el que hizo churn
        active_assets = active_by_org.get(key, [])
        active_assets = [a for a in active_assets if clean(a.get('Assets Asset Name', '')) != asset_name]

        # Sumar MRR de assets activos
        active_mrr = sum(money(a.get('Assets Current MRR', 0)) for a in active_assets)

        previous_mrr = money(row.get('Assets Previous MRR', 0))
        current_mrr = money(row.get('Assets Current MRR', 0))
        operation = clean(row.get('Assets Operation', ''))

        # Filtro: previousMrr > 0 OR currentMrr == 0 OR operation == 'cancellation'
        if not (previous_mrr > 0 or current_mrr == 0 or operation == 'cancellation'):
            continue

        # applies = no hay otros assets activos
        applies = len(active_assets) == 0

        results.append({
            'applies': applies,
            'previous_mrr': previous_mrr,
            'current_mrr': current_mrr,
            'org_id': row.get('Organizations Organization ID', ''),
            'organization': row.get('Organizations Organization Name', ''),
            'asset_name': row.get('Assets Asset Name', ''),
            'operation': row.get('Assets Operation', ''),
            'assigned_to': row.get('Organizations Assigned To', ''),
            'tier': row.get('Organizations Freshworks Tier', 'Sin dato'),
            'active_assets': active_assets,
            'active_mrr': active_mrr,
            'raw_row': row  # Guardamos la fila original para ref
        })

    # Ordenar: primero applies=True, luego por previousMrr descendente
    results.sort(key=lambda r: (-int(r['applies']), -r['previous_mrr']))

    return results

def main():
    # Rutas de los Excel
    descargas = DOWNLOADS

    # Buscar los archivos más recientes
    churn_pattern = 'Mid Market Metricas 2025 - ChurnAssets_*.xlsx'
    oa_pattern = 'Mid Market Metricas 2025 - OA_*.xlsx'

    churn_files = sorted(descargas.glob(churn_pattern), key=lambda p: p.stat().st_mtime, reverse=True)
    oa_files = sorted(descargas.glob(oa_pattern), key=lambda p: p.stat().st_mtime, reverse=True)

    if not churn_files:
        print(f"ERROR: No se encontró {churn_pattern} en {descargas}")
        sys.exit(1)
    if not oa_files:
        print(f"ERROR: No se encontró {oa_pattern} en {descargas}")
        sys.exit(1)

    churn_file = churn_files[0]
    oa_file = oa_files[0]

    print(f"Cargando ChurnAssets: {churn_file.name}")
    print(f"Cargando OA: {oa_file.name}")

    # Cargar datos
    churn_rows = load_excel_as_dicts(str(churn_file))
    oa_rows = load_excel_as_dicts(str(oa_file))

    print(f"  ChurnAssets: {len(churn_rows)} filas")
    print(f"  OA: {len(oa_rows)} filas")

    # Aplicar lógica de validación
    validated = churn_validation_rows(churn_rows, oa_rows)

    # Filtrar solo las que aplican
    applies = [r for r in validated if r['applies']]

    print(f"\nResultados:")
    print(f"  Total filas después de filtros: {len(validated)}")
    print(f"  Churn aplicable (applies=True): {len(applies)}")

    # Guardar para el siguiente paso
    output_file = CHURN_BATCH_DIR / 'churn_aplicable.json'
    output_file.parent.mkdir(parents=True, exist_ok=True)

    # Serializar: quitar 'raw_row' y 'active_assets' (objetos no serializables)
    applies_clean = []
    for r in applies:
        applies_clean.append({
            'org_id': r['org_id'],
            'organization': r['organization'],
            'asset_name': r['asset_name'],
            'assigned_to': r['assigned_to'],
            'tier': r['tier'],
            'previous_mrr': r['previous_mrr'],
            'current_mrr': r['current_mrr'],
            'operation': r['operation'],
            'active_mrr': r['active_mrr'],
            'active_assets_count': len(r['active_assets'])
        })

    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(applies_clean, f, indent=2, ensure_ascii=False)

    print(f"\nLista de churn aplicable guardada en: {output_file}")
    print(f"\nCuentas con churn aplicable ({len(applies)}):")
    print("-" * 100)
    for r in applies:
        print(f"{r['org_id']:10} | {r['organization']:35} | ${r['previous_mrr']:7.2f} MRR | {r['assigned_to']}")

if __name__ == '__main__':
    main()
