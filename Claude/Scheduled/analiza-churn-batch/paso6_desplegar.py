#!/usr/bin/env python3
"""Los dos extremos del despliegue de churn a SharePoint, que no son la subida.

El 2-oct-2026 el despliegue subio 4 de 5 informes y dejo el Excel sin tocar: la
subida de CIMA (el HTML mas grande) fallo, el agente se quedo peleando con ella
y la corrida murio antes de llegar a actualizar los enlaces. Las 4 subidas que
si habian salido bien no sirvieron de nada, porque el Excel seguia apuntando a
rutas locales.

La subida tiene que pasar por el conector de Microsoft 365 (no hay credenciales
de Graph en esta maquina, ver sales-assistant/backend/.env.example), asi que esa
parte la hace el agente. Lo que NO tiene por que depender de que el agente
sobreviva es leer el Excel y escribir los enlaces: eso es determinista y vive
aqui.

    python paso6_desplegar.py --preparar
        Encuentra el resumen mas reciente y deja despliegue_items.json con la
        lista de informes a subir.

    python paso6_desplegar.py --aplicar
        Lee despliegue_resultado.json (lo escribe el agente con el webUrl de
        cada subida) y actualiza los enlaces del Excel, una sola vez.

Si una subida fallo, su fila conserva el enlace local a proposito, para que se
vea cual quedo pendiente. Y como --aplicar es idempotente, si la corrida se
muere despues de subir basta con volver a correrlo: no hay que subir de nuevo.
"""
import argparse
import json
import re
import sys
from urllib.parse import quote, unquote, urlsplit, urlunsplit
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "scripts"))
from paths import CHURN_BATCH_DIR, CHURN_INFORMES_DIR

ITEMS_FILE = CHURN_BATCH_DIR / "despliegue_items.json"
RESULTADO_FILE = CHURN_BATCH_DIR / "despliegue_resultado.json"

# Acepta el sufijo -N que deja el paso 5 cuando el Excel del dia estaba abierto
# en Excel. Sin esto, --preparar podia agarrar el resumen viejo y dejar el bueno
# sin actualizar.
RESUMEN_RE = re.compile(r"^resumen-churn-(\d{4}-\d{2}-\d{2})(?:-(\d+))?\.xlsx$")

COL_INFORME_HEADER = "Informe"


def resumen_mas_reciente():
    """El resumen de fecha mas alta y, dentro de esa fecha, el sufijo mas alto.

    Se ordena por (fecha, sufijo) y no por mtime: el mtime cambia con solo
    abrir y guardar el archivo en Excel, y eso no lo vuelve el mas nuevo."""
    candidatos = []
    for p in CHURN_INFORMES_DIR.iterdir():
        m = RESUMEN_RE.match(p.name)
        if m:
            candidatos.append((m.group(1), int(m.group(2) or 0), p))
    if not candidatos:
        return None
    return max(candidatos)[2]


def normalizar_url(url):
    """Codifica los espacios del path y deja igual lo que ya venia codificado.

    El conector devuelve el webUrl con %20 unas veces y con espacios literales
    otras, segun por donde se consulte. Un enlace con espacios sin codificar en
    una celda de Excel es fragil, asi que se normalizan todos aqui y no depende
    de como lo reporte cada subida (2-oct-2026)."""
    partes = urlsplit(url)
    return urlunsplit((
        partes.scheme,
        partes.netloc,
        quote(unquote(partes.path), safe="/"),
        partes.query,
        partes.fragment,
    ))


def columna_informe(ws):
    """Posicion de la columna 'Informe' leyendo la cabecera.

    Se busca por nombre y no por indice fijo porque el resumen paso de 12 a 15
    columnas el 1-oct-2026, y un indice escrito a mano habria apuntado a la
    columna equivocada sin avisar."""
    for celda in ws[1]:
        if (celda.value or "").strip().lower() == COL_INFORME_HEADER.lower():
            return celda.column
    raise SystemExit(f"ERROR: el Excel no tiene columna '{COL_INFORME_HEADER}'.")


def preparar():
    xlsx = resumen_mas_reciente()
    if not xlsx:
        raise SystemExit(f"ERROR: no hay ningun resumen-churn-*.xlsx en {CHURN_INFORMES_DIR}")
    print(f"Resumen mas reciente: {xlsx.name}")

    wb = load_workbook(xlsx)
    ws = wb.active
    col = columna_informe(ws)

    items, omitidos = [], []
    for fila in range(2, ws.max_row + 1):
        org = ws.cell(fila, 1).value
        if not org:
            continue
        celda = ws.cell(fila, col)
        destino = celda.hyperlink.target if celda.hyperlink else None
        if not destino:
            omitidos.append((org, "sin enlace"))
            continue
        if not destino.lower().startswith("file:"):
            omitidos.append((org, "ya desplegado"))
            continue

        # file:///C:\ruta  ->  C:\ruta
        ruta = Path(destino[len("file:///"):].replace("/", "\\"))
        if not ruta.exists():
            omitidos.append((org, f"no existe {ruta.name}"))
            continue
        items.append({
            "org": org,
            "fila": fila,
            "nombre_html": ruta.name,
            "ruta_local": str(ruta),
            "bytes": ruta.stat().st_size,
        })

    ITEMS_FILE.write_text(
        json.dumps({"xlsx": str(xlsx), "items": items}, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print(f"\nInformes a subir ({len(items)}):")
    for it in items:
        print(f"  {it['org']:<30} {it['nombre_html']}  ({it['bytes']:,} bytes)")
    for org, motivo in omitidos:
        print(f"  OMITIDA: {org} ({motivo})")
    print(f"\nLista guardada en: {ITEMS_FILE}")
    print(f"Cuando subas, deja los resultados en: {RESULTADO_FILE}")
    print('  formato: [{"org": "...", "success": true, "webUrl": "..."}, ...]')


def aplicar():
    if not ITEMS_FILE.exists():
        raise SystemExit(f"ERROR: falta {ITEMS_FILE}. Corre antes --preparar.")
    if not RESULTADO_FILE.exists():
        raise SystemExit(
            f"ERROR: falta {RESULTADO_FILE}. Ahi van los resultados de las subidas."
        )

    datos = json.loads(ITEMS_FILE.read_text(encoding="utf-8"))
    xlsx = Path(datos["xlsx"])
    resultados = {r["org"]: r for r in json.loads(RESULTADO_FILE.read_text(encoding="utf-8"))}

    wb = load_workbook(xlsx)
    ws = wb.active
    col = columna_informe(ws)

    actualizadas, pendientes = [], []
    for it in datos["items"]:
        res = resultados.get(it["org"])
        if not res or not res.get("success") or not res.get("webUrl"):
            # Se deja el enlace local a proposito: asi se ve cual falta.
            pendientes.append((it["org"], (res or {}).get("error", "sin resultado")))
            continue
        celda = ws.cell(it["fila"], col)
        celda.hyperlink = normalizar_url(res["webUrl"])
        celda.value = "Ver informe"
        celda.font = Font(color="0563C1", underline="single")
        actualizadas.append(it["org"])

    # Mismo criterio que el paso 5 y que gbs_deal_checker: si el Excel esta
    # abierto se guarda al lado en vez de perder el trabajo.
    try:
        wb.save(xlsx)
        destino = xlsx
    except PermissionError:
        # Se conserva la FECHA DEL ORIGEN y solo se sube el sufijo. Con la fecha
        # de hoy, desplegar el 2-oct un resumen del 1-oct dejaba un
        # resumen-churn-2026-10-02-2.xlsx que parecia una corrida nueva de ese
        # dia y no lo que es: una copia del resumen del 1-oct con los enlaces
        # ya puestos (2-oct-2026).
        fecha = RESUMEN_RE.match(xlsx.name).group(1)
        n = 2
        while (CHURN_INFORMES_DIR / f"resumen-churn-{fecha}-{n}.xlsx").exists():
            n += 1
        destino = CHURN_INFORMES_DIR / f"resumen-churn-{fecha}-{n}.xlsx"
        wb.save(destino)
        print(f"AVISO: '{xlsx.name}' esta abierto en Excel y no se pudo sobrescribir.")
        print(f"       Se guardo en '{destino.name}'. Cerra el archivo y renombralo.")

    print(f"Enlaces actualizados ({len(actualizadas)}):")
    for org in actualizadas:
        print(f"  {org}")
    if pendientes:
        print(f"\nSiguen con enlace local ({len(pendientes)}):")
        for org, motivo in pendientes:
            print(f"  {org}: {motivo}")
    print(f"\nGuardado: {destino}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--preparar", action="store_true")
    g.add_argument("--aplicar", action="store_true")
    args = ap.parse_args()
    preparar() if args.preparar else aplicar()
