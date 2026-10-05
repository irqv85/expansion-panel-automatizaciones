"""
Genera un Excel local del plan de compensacion MID Growth, con una hoja por
vendedor lista para copiar y pegar en la hoja oficial de SharePoint.

Lee el reporte DIQ que ya descargas en la carpeta de Descargas (el mas
reciente). No consulta vTiger por internet.

Solo llena las columnas que van llenas en las pestañas por vendedor:
    Opportunity Name | Deal ID | ECD | Comision Rate | MRR
La columna Domain y TODO el bloque FINANCE VERIFICATION se dejan vacios --
esos los llena Finanzas.

Regla del trimestre:
  - Sales Stage = Closed Won
  - Vendedor del equipo Mid/SMB (Vendedor 4 (alias) excluido)
  - Expected Close Date dentro del trimestre
  - Clasificacion de la comision:
      ProfessionalServices GB/Vendor -> "100%"      (suma a sip+ssi)
      cualquier otro tipo            -> "mrr 100%"  (suma a MRR 100%)

Uso:
  python "Plan Compensacion.py"                  Trimestre actual
  python "Plan Compensacion.py" --trimestre 2026Q3
  python "Plan Compensacion.py" --desde 2026-07-01 --hasta 2026-09-30
"""
import argparse
import sys
import traceback
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---- Configuracion ----
from paths import DOWNLOADS, COMP_PLAN_DIR as OUTPUT_DIR

# Nombre en el reporte -> nombre para la hoja. Vendedor 4 (alias) queda fuera
# a proposito.
# nombre del vendedor -> nombre de su hoja, desde equipo.json
from equipo import REPS

# Columnas del reporte DIQ.
COL_DEAL_ID = ["Deals Deal ID", "Deals Deal Id", "Deal ID"]
COL_NOMBRE = ["Deals Deal Name", "Deal Name"]
COL_ECD = ["Deals Expected Close Date", "Expected Close Date"]
COL_MONTO = ["Deals Amount", "Amount"]
COL_ETAPA = ["Deals Sales Stage", "Sales Stage"]
COL_VENDEDOR = ["Deals Assigned To", "Assigned To"]
COL_TIPO = ["Deals New/Renew/ProServ", "New/Renew/ProServ"]
COL_ORG_DEAL = ["Deals Organization Name", "Organization Name"]

# Booster Hours, SSI, HRS y demas servicios cuentan como professional
# services, sean de GB o de vendor: todos van al plan con rate "100%".
TIPOS_SERVICIO = {"professionalservices gb", "professionalservices vendor", "ssi add on"}

FONT_NAME = "Arial"
INK = "12112C"
HEADER_FILL = PatternFill("solid", fgColor="25235A")
FINANCE_FILL = PatternFill("solid", fgColor="F1F3F5")
THIN = Side(style="thin", color="DEE2E6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def log(msg):
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


def quarter_bounds(ref=None):
    ref = ref or date.today()
    q = (ref.month - 1) // 3 + 1
    return _bounds(ref.year, q)


def _bounds(year, q):
    start = date(year, 3 * (q - 1) + 1, 1)
    end = date(year + (1 if q == 4 else 0), 1 if q == 4 else 3 * q + 1, 1) - timedelta(days=1)
    return start, end, f"{year}Q{q}"


def parse_quarter(text):
    return _bounds(int(text[:4]), int(text[-1]))


def find_reporte(patron, descripcion, obligatorio=True):
    archivos = sorted(DOWNLOADS.glob(patron), key=lambda p: p.stat().st_mtime)
    if not archivos:
        if obligatorio:
            raise FileNotFoundError(
                f"No se encontro el reporte {descripcion} ({patron}) en {DOWNLOADS}. "
                f"Descargalo de vTiger primero."
            )
        return None
    return archivos[-1]


def find_diq():
    return find_reporte("*DIQ*.xlsx", "DIQ")


def leer_reporte(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    # Un export de vTiger sin resultados viene con CERO filas, ni siquiera la
    # cabecera, y next() reventaba con StopIteration. Como StopIteration no
    # lleva mensaje, el manejador de abajo imprimia "ERROR: " a secas y no
    # habia forma de saber que habia pasado. Caso real: el 1-oct-2026, primer
    # dia de Q4, el reporte de ChurnAssets salio vacio porque todavia no habia
    # churn en el trimestre, y el plan de compensacion fallaba sin explicar
    # nada (1-oct-2026).
    primera = next(it, None)
    if primera is None:
        wb.close()
        return [], []
    header = [str(h).strip() if h is not None else "" for h in primera]
    filas = [dict(zip(header, r)) for r in it if any(v not in (None, "") for v in r)]
    wb.close()
    return header, filas


def buscar_columna(header, candidatas):
    normal = {h.strip().lower(): h for h in header}
    for c in candidatas:
        if c.strip().lower() in normal:
            return normal[c.strip().lower()]
    return None


def a_fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    if not texto:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(texto[:10], fmt).date()
        except ValueError:
            continue
    return None


def a_monto(valor):
    if valor in (None, ""):
        return 0.0
    try:
        return float(str(valor).replace("$", "").replace(",", ""))
    except ValueError:
        return 0.0


def clasificar(desde, hasta, header, filas):
    """Devuelve los deals del trimestre agrupados por vendedor."""
    c_id = buscar_columna(header, COL_DEAL_ID)
    c_nombre = buscar_columna(header, COL_NOMBRE)
    c_ecd = buscar_columna(header, COL_ECD)
    c_monto = buscar_columna(header, COL_MONTO)
    c_etapa = buscar_columna(header, COL_ETAPA)
    c_vend = buscar_columna(header, COL_VENDEDOR)
    c_tipo = buscar_columna(header, COL_TIPO)

    faltan = [n for n, c in (("Deal ID", c_id), ("Deal Name", c_nombre),
                             ("Expected Close Date", c_ecd), ("Amount", c_monto),
                             ("Sales Stage", c_etapa), ("Assigned To", c_vend),
                             ("New/Renew/ProServ", c_tipo)) if c is None]
    if faltan:
        raise RuntimeError(
            "Al reporte DIQ le faltan estas columnas: " + ", ".join(faltan)
        )

    por_vendedor = {nombre: [] for nombre in REPS.values()}
    vistos = {nombre: set() for nombre in REPS.values()}

    for r in filas:
        if str(r.get(c_etapa) or "").strip() != "Closed Won":
            continue
        rep = REPS.get(str(r.get(c_vend) or "").strip())
        if rep is None:
            continue

        ecd = a_fecha(r.get(c_ecd))
        if ecd is None or not (desde <= ecd <= hasta):
            continue

        deal_id = str(r.get(c_id) or "").strip()
        # DIQ trae "Products Product Name": un deal con varios productos
        # asociados sale en una fila por producto, con el mismo monto
        # repetido en cada una. Sin este filtro, ese deal se suma dos o mas
        # veces al plan de compensacion.
        if deal_id and deal_id in vistos[rep]:
            continue
        if deal_id:
            vistos[rep].add(deal_id)

        tipo = str(r.get(c_tipo) or "").strip()
        es_servicio = tipo.lower() in TIPOS_SERVICIO
        nombre = str(r.get(c_nombre) or "").strip()
        monto = a_monto(r.get(c_monto))

        por_vendedor[rep].append({
            "nombre": nombre,
            "deal_id": deal_id,
            "ecd": f"{ecd:%d/%m/%Y}" if ecd else "",
            "rate": "100%" if es_servicio else "mrr 100%",
            "mrr": monto,
        })

    for lista in por_vendedor.values():
        lista.sort(key=lambda f: (a_fecha(f["ecd"]) or date.min, f["deal_id"]))
    return por_vendedor


# ---------- Excel ----------

COLUMNAS = [
    ("Opportunity Name", True),
    ("Deal ID", True),
    ("ECD", True),
    ("Domain", False),
    ("Comision Rate", True),
    ("MRR", True),
    ("Payment to GB?", False),
    ("Partner Portal Check", False),
    ("Freshworks ID", False),
    ("Freshwork Invoice", False),
    ("GB INVOICE", False),
    ("Payment date", False),
    ("Amount", False),
    ("Cliente pago", False),
    ("COMM GB", False),
    ("Pendiente para el siguiente Q", False),
]


def _titulo(ws, texto, ncols):
    ws["A1"] = texto
    ws["A1"].font = Font(name=FONT_NAME, size=13, bold=True, color=INK)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 2))


def escribir_hoja_vendedor(wb, rep, filas):
    ws = wb.create_sheet(rep[:31])
    _titulo(ws, rep, len(COLUMNAS))

    ws["A2"] = ("Copia y pega solo las columnas con datos. Domain y el bloque de "
                "FINANCE VERIFICATION se dejan vacios a proposito: los llena Finanzas.")
    ws["A2"].font = Font(name=FONT_NAME, size=9, italic=True, color="5B5F6B")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUMNAS))

    enc = 4
    for i, (titulo, se_llena) in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=enc, column=i, value=titulo)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = BORDER
        c.font = Font(name=FONT_NAME, size=9, bold=True,
                      color="FFFFFF" if se_llena else "C9C7E0")

    primera = enc + 1
    for j, f in enumerate(filas):
        r = primera + j
        ws.cell(row=r, column=1, value=f["nombre"])
        ws.cell(row=r, column=2, value=f["deal_id"])
        ws.cell(row=r, column=3, value=f["ecd"])
        ws.cell(row=r, column=5, value=f["rate"])
        c_mrr = ws.cell(row=r, column=6, value=f["mrr"])
        c_mrr.number_format = '$#,##0.00;($#,##0.00);-'
        for i in range(1, len(COLUMNAS) + 1):
            celda = ws.cell(row=r, column=i)
            celda.font = Font(name=FONT_NAME, size=10, color=INK)
            celda.border = BORDER
            if not COLUMNAS[i - 1][1]:
                celda.fill = FINANCE_FILL

    ultima = primera + len(filas) - 1 if filas else None
    base = (ultima if ultima else enc) + 2
    ws.cell(row=base, column=1, value="Commssion Rate").font = Font(name=FONT_NAME, size=10, bold=True, color=INK)
    ws.cell(row=base, column=2, value="MRR").font = Font(name=FONT_NAME, size=10, bold=True, color=INK)

    if ultima:
        rr, rm = f"$E${primera}:$E${ultima}", f"$F${primera}:$F${ultima}"
        f_sip, f_lic = f'=SUMIF({rr},"100%",{rm})', f'=SUMIF({rr},"mrr 100%",{rm})'
    else:
        f_sip = f_lic = 0

    for k, (etiqueta, valor) in enumerate(
        [("sip+ssi 100.00%", f_sip), ("MRR 100%", f_lic), ("Total", f"=B{base + 1}+B{base + 2}")],
        start=1,
    ):
        r = base + k
        negrita = etiqueta == "Total"
        c1 = ws.cell(row=r, column=1, value=etiqueta)
        c2 = ws.cell(row=r, column=2, value=valor)
        c1.font = c2.font = Font(name=FONT_NAME, size=10, bold=negrita, color=INK)
        c2.number_format = '$#,##0.00;($#,##0.00);-'
        c1.border = c2.border = BORDER

    for i, ancho in enumerate([46, 12, 12, 14, 14, 13] + [16] * (len(COLUMNAS) - 6), start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = ws.cell(row=primera, column=1)
    return ws, primera, ultima


def _org_key(row):
    """Misma llave de organizacion que usa el Dashboard: Org ID si existe,
    si no el nombre."""
    oid = str(row.get("Organizations Organization ID") or "").strip().lower()
    if oid:
        return f"id:{oid}"
    nombre = str(row.get("Organizations Organization Name") or "").strip().lower()
    return f"name:{nombre}" if nombre else ""


COL_CHURN_FECHA = ("Assets VT Change Date", "Assets Change Date", "VT Change Date")


def churn_del_trimestre(churn_filas, desde, hasta):
    """Deja solo los assets que cambiaron DENTRO del trimestre que se pide.

    El export de ChurnAssets trae todo lo que vTiger tenga cargado, sin acotar
    al periodo. Hasta el 1-oct-2026 el plan lo usaba entero, con lo que un plan
    de Q4 listaba los 9 churns de Q3 (del 24-jul al 10-sep) como si fueran del
    trimestre en curso. En un plan de compensacion eso significa descontarle a
    un vendedor un churn que ya se conto en el trimestre anterior.

    Una fila sin fecha legible se conserva: es preferible que salga y se revise
    a mano antes que desaparecer sin que nadie se entere."""
    if not churn_filas:
        return [], 0
    columna = buscar_columna(list(churn_filas[0].keys()), COL_CHURN_FECHA)
    if not columna:
        return churn_filas, 0  # sin columna de fecha no se puede acotar
    dentro, fuera = [], 0
    for r in churn_filas:
        f = a_fecha(r.get(columna))
        if f is None or (desde <= f <= hasta):
            dentro.append(r)
        else:
            fuera += 1
    return dentro, fuera


def analizar_churn(churn_filas, oa_filas, diq_header, diq_filas):
    """Replica la validacion de churn del Dashboard y agrega el analisis de
    si el churn le corresponde al vendedor.

    Regla del Dashboard: el churn de un asset solo aplica si en OA no queda
    NINGUN otro asset activo (Current MRR > 0) de la misma organizacion. Si
    queda algo activo es churn parcial y no cuenta como churn de la cuenta.

    Sobre la responsabilidad, las notas del plan dicen que el churn le toca al
    AM que gestiono la cuenta ("al menos una venta del AM", o mas de 3 meses de
    gestion). Con los reportes locales se puede verificar la venta (Closed Won
    del AM en esa organizacion); los 3 meses de gestion no, porque ningun
    reporte trae la fecha en que se le asigno la cuenta -- esos casos salen
    marcados como "Revisar"."""
    activos_por_org = {}
    for a in oa_filas:
        if a_monto(a.get("Assets Current MRR")) <= 0:
            continue
        k = _org_key(a)
        if k:
            activos_por_org.setdefault(k, []).append(a)

    # Ventas Closed Won por (organizacion, vendedor), para el criterio de
    # "al menos una venta del AM".
    c_org = buscar_columna(diq_header, COL_ORG_DEAL)
    c_etapa = buscar_columna(diq_header, COL_ETAPA)
    c_vend = buscar_columna(diq_header, COL_VENDEDOR)
    c_id = buscar_columna(diq_header, COL_DEAL_ID)
    ventas = {}
    vistos_deal_id = set()
    if c_org and c_etapa and c_vend:
        for r in diq_filas:
            if str(r.get(c_etapa) or "").strip() != "Closed Won":
                continue
            # Igual que en clasificar(): un deal con varios productos sale
            # repetido en DIQ, y no debe contarse como varias ventas.
            deal_id = str(r.get(c_id) or "").strip() if c_id else ""
            if deal_id:
                if deal_id in vistos_deal_id:
                    continue
                vistos_deal_id.add(deal_id)
            clave = (str(r.get(c_org) or "").strip().lower(),
                     str(r.get(c_vend) or "").strip().lower())
            ventas[clave] = ventas.get(clave, 0) + 1

    resultado = []
    for row in churn_filas:
        previo = a_monto(row.get("Assets Previous MRR"))
        actual = a_monto(row.get("Assets Current MRR"))
        operacion = str(row.get("Assets Operation") or "").strip()

        # Mismo filtro que churnValidationRows del Dashboard.
        if not (previo > 0 or actual == 0 or operacion.lower() == "cancellation"):
            continue

        asset = str(row.get("Assets Asset Name") or "").strip()
        k = _org_key(row)
        otros = [a for a in activos_por_org.get(k, [])
                 if str(a.get("Assets Asset Name") or "").strip().lower() != asset.lower()]
        mrr_activo = sum(a_monto(a.get("Assets Current MRR")) for a in otros)
        aplica = len(otros) == 0
        perdido = max(0.0, previo - actual)

        vendedor = str(row.get("Organizations Assigned To") or "").strip()
        organizacion = str(row.get("Organizations Organization Name") or "").strip()
        es_del_equipo = vendedor in REPS
        n_ventas = ventas.get((organizacion.lower(), vendedor.lower()), 0)

        if not es_del_equipo:
            corresponde = "No"
            motivo = (f"La cuenta esta asignada a {vendedor or 'nadie'}, que no es del "
                      f"equipo Mid/SMB de este plan.")
        elif not aplica:
            corresponde = "No"
            motivo = (f"Churn parcial: la organizacion sigue con {len(otros)} asset(s) "
                      f"activo(s) por ${mrr_activo:,.2f} de MRR. Segun el Dashboard el "
                      f"churn solo aplica si no queda ningun asset activo.")
        elif perdido <= 0:
            corresponde = "No"
            motivo = (f"No hubo perdida de MRR (anterior ${previo:,.2f} vs actual "
                      f"${actual:,.2f}), asi que no hay churn que asignar.")
        elif n_ventas > 0:
            corresponde = "Si"
            motivo = (f"Churn total y el AM tiene {n_ventas} venta(s) Closed Won en esta "
                      f"cuenta, que es el criterio del plan.")
        else:
            corresponde = "Revisar"
            motivo = ("Churn total, pero en el reporte no aparece ninguna venta Closed Won "
                      "del AM en esta cuenta. Valida si la gestiono mas de 3 meses: esa "
                      "fecha no viene en los reportes.")

        resultado.append({
            "vendedor": vendedor,
            "organizacion": organizacion,
            "org_id": str(row.get("Organizations Organization ID") or "").strip(),
            "asset": asset,
            "operacion": operacion,
            "tier": str(row.get("Organizations Freshworks Tier") or "Sin dato").strip(),
            "previo": previo,
            "actual": actual,
            "perdido": perdido,
            "aplica": "Si" if aplica else "No",
            "n_otros": len(otros),
            "mrr_activo": mrr_activo,
            "otros": ", ".join(str(a.get("Assets Asset Name") or "") for a in otros)[:120],
            "corresponde": corresponde,
            "motivo": motivo,
        })

    orden = {"Si": 0, "Revisar": 1, "No": 2}
    resultado.sort(key=lambda r: (orden.get(r["corresponde"], 3), -r["perdido"]))
    return resultado


COLUMNAS_CHURN = [
    ("Vendedor", 22), ("Organizacion", 30), ("Org ID", 11), ("Asset con churn", 26),
    ("Operacion", 14), ("Tier", 12), ("Previous MRR", 14), ("Current MRR", 14),
    ("MRR perdido", 14), ("Aplica churn", 13), ("Assets activos OA", 11),
    ("MRR activo OA", 14), ("Otros assets activos", 34),
    ("Le corresponde al vendedor?", 15), ("Motivo", 80),
]


def escribir_churn(wb, filas_churn, reporte_churn, reporte_oa):
    ws = wb.create_sheet("Churn")
    _titulo(ws, "Churn - analisis de responsabilidad", len(COLUMNAS_CHURN))
    fuentes = f"{reporte_churn.name} vs {reporte_oa.name}" if reporte_churn and reporte_oa else "sin reportes"
    ws["A2"] = ("Regla del Dashboard: el churn aplica solo si no queda ningun otro asset activo "
                f"en OA para la misma organizacion. Fuentes: {fuentes}.")
    ws["A2"].font = Font(name=FONT_NAME, size=9, italic=True, color="5B5F6B")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUMNAS_CHURN))

    for i, (titulo, _) in enumerate(COLUMNAS_CHURN, start=1):
        c = ws.cell(row=4, column=i, value=titulo)
        c.font = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = BORDER

    VERDE = PatternFill("solid", fgColor="E4F3EA")
    AMARILLO = PatternFill("solid", fgColor="FFF4CE")

    if not filas_churn:
        ws.cell(row=5, column=1, value="Sin churn en el reporte.").font = Font(
            name=FONT_NAME, size=10, color=INK)

    for j, f in enumerate(filas_churn):
        r = 5 + j
        valores = [
            f["vendedor"], f["organizacion"], f["org_id"], f["asset"], f["operacion"],
            f["tier"], f["previo"], f["actual"], f["perdido"], f["aplica"],
            f["n_otros"], f["mrr_activo"], f["otros"], f["corresponde"], f["motivo"],
        ]
        for i, v in enumerate(valores, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=FONT_NAME, size=10, color=INK,
                          bold=(i == 14))
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=(i in (13, 15)))
            if i in (7, 8, 9, 12):
                c.number_format = '$#,##0.00;($#,##0.00);-'
            if f["corresponde"] == "Si":
                c.fill = VERDE
            elif f["corresponde"] == "Revisar":
                c.fill = AMARILLO

    fila_total = 5 + max(len(filas_churn), 1) + 1
    ws.cell(row=fila_total, column=8, value="Churn que si corresponde:").font = Font(
        name=FONT_NAME, size=10, bold=True, color=INK)
    if filas_churn:
        primera, ultima = 5, 5 + len(filas_churn) - 1
        formula = f'=SUMIF($N${primera}:$N${ultima},"Si",$I${primera}:$I${ultima})'
    else:
        formula = 0
    c = ws.cell(row=fila_total, column=9, value=formula)
    c.font = Font(name=FONT_NAME, size=10, bold=True, color=INK)
    c.number_format = '$#,##0.00;($#,##0.00);-'
    c.border = BORDER

    for i, (_, ancho) in enumerate(COLUMNAS_CHURN, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = ws.cell(row=5, column=1)


def escribir_resumen(wb, etiqueta, desde, hasta, por_vendedor, hojas, reporte):
    ws = wb.create_sheet("Resumen", 0)
    _titulo(ws, f"Plan de compensacion MID Growth - {etiqueta}", 5)
    ws["A2"] = (f"Closed Won con Expected Close Date del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}. "
                f"Fuente: {reporte.name}. Generado {datetime.now():%d/%m/%Y %H:%M}. "
                f"Vendedor 4 (alias) excluido.")
    ws["A2"].font = Font(name=FONT_NAME, size=9, italic=True, color="5B5F6B")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)

    for i, t in enumerate(["Vendedor", "Deals", "sip+ssi (100%)", "MRR (mrr 100%)", "Total"], start=1):
        c = ws.cell(row=4, column=i, value=t)
        c.font = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.border = BORDER

    r = 5
    for rep in REPS.values():
        filas = por_vendedor.get(rep, [])
        hoja, primera, ultima = hojas[rep]
        ws.cell(row=r, column=1, value=rep)
        ws.cell(row=r, column=2, value=len(filas))
        if ultima:
            ref = f"'{hoja.title}'"
            ws.cell(row=r, column=3, value=f'=SUMIF({ref}!$E${primera}:$E${ultima},"100%",{ref}!$F${primera}:$F${ultima})')
            ws.cell(row=r, column=4, value=f'=SUMIF({ref}!$E${primera}:$E${ultima},"mrr 100%",{ref}!$F${primera}:$F${ultima})')
        else:
            ws.cell(row=r, column=3, value=0)
            ws.cell(row=r, column=4, value=0)
        ws.cell(row=r, column=5, value=f"=C{r}+D{r}")
        for i in range(1, 6):
            c = ws.cell(row=r, column=i)
            c.font = Font(name=FONT_NAME, size=10, color=INK)
            c.border = BORDER
            if i >= 3:
                c.number_format = '$#,##0.00;($#,##0.00);-'
        r += 1

    ws.cell(row=r, column=1, value="Total equipo").font = Font(name=FONT_NAME, size=10, bold=True, color=INK)
    for i, letra in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
        c = ws.cell(row=r, column=i, value=f"=SUM({letra}5:{letra}{r - 1})")
        c.font = Font(name=FONT_NAME, size=10, bold=True, color=INK)
        c.border = BORDER
        if i >= 3:
            c.number_format = '$#,##0.00;($#,##0.00);-'

    for i, ancho in enumerate([26, 10, 18, 18, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def main():
    parser = argparse.ArgumentParser(description="Genera el plan de compensacion local desde el reporte DIQ.")
    parser.add_argument("--trimestre", help="Trimestre a generar, ej. 2026Q3. Por defecto el actual.")
    parser.add_argument("--desde", help="Fecha inicial YYYY-MM-DD (tiene prioridad sobre --trimestre).")
    parser.add_argument("--hasta", help="Fecha final YYYY-MM-DD.")
    args = parser.parse_args()

    if args.desde and args.hasta:
        desde = datetime.strptime(args.desde, "%Y-%m-%d").date()
        hasta = datetime.strptime(args.hasta, "%Y-%m-%d").date()
        etiqueta = f"{desde:%d-%m-%Y} a {hasta:%d-%m-%Y}"
    elif args.trimestre:
        desde, hasta, etiqueta = parse_quarter(args.trimestre)
    else:
        desde, hasta, etiqueta = quarter_bounds()

    reporte = find_diq()
    edad = datetime.now() - datetime.fromtimestamp(reporte.stat().st_mtime)
    log(f"Reporte: {reporte.name}")
    if edad.days >= 1:
        log(f"AVISO: el reporte tiene {edad.days} dia(s). Descargalo de nuevo para "
            f"incluir los deals cerrados desde entonces.")
    log(f"Trimestre: {etiqueta} ({desde} a {hasta})")

    header, filas = leer_reporte(reporte)
    log(f"Filas leidas del reporte: {len(filas)}")

    por_vendedor = clasificar(desde, hasta, header, filas)

    for rep, lista in por_vendedor.items():
        total = sum(f["mrr"] for f in lista)
        log(f"  {rep}: {len(lista)} deal(s), MRR total ${total:,.2f}")

    # --- Churn ---
    reporte_churn = find_reporte("*ChurnAssets*.xlsx", "ChurnAssets", obligatorio=False)
    reporte_oa = find_reporte("*- OA_*.xlsx", "OA", obligatorio=False)
    filas_churn = []
    if reporte_churn and reporte_oa:
        _, churn_filas = leer_reporte(reporte_churn)
        _, oa_filas = leer_reporte(reporte_oa)
        churn_filas, churn_fuera = churn_del_trimestre(churn_filas, desde, hasta)
        if churn_fuera:
            log(f"Churn: {churn_fuera} asset(s) fuera de {etiqueta}, no se cuentan "
                f"(su fecha de cambio cae en otro trimestre).")
        filas_churn = analizar_churn(churn_filas, oa_filas, header, filas)
        n_si = sum(1 for f in filas_churn if f["corresponde"] == "Si")
        n_rev = sum(1 for f in filas_churn if f["corresponde"] == "Revisar")
        perdido = sum(f["perdido"] for f in filas_churn if f["corresponde"] == "Si")
        log(f"Churn: {len(filas_churn)} asset(s) revisados | corresponde al vendedor: {n_si} "
            f"(${perdido:,.2f}) | por revisar: {n_rev}")
    else:
        log("AVISO: falta el reporte ChurnAssets o el de OA; la pestaña Churn queda vacia.")

    wb = Workbook()
    wb.remove(wb.active)
    hojas = {rep: escribir_hoja_vendedor(wb, rep, por_vendedor[rep]) for rep in REPS.values()}
    escribir_churn(wb, filas_churn, reporte_churn, reporte_oa)
    escribir_resumen(wb, etiqueta, desde, hasta, por_vendedor, hojas, reporte)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    salida = OUTPUT_DIR / f"Plan Compensacion MID Growth {etiqueta}.xlsx"
    wb.save(salida)
    log(f"Listo: {salida}")
    return salida


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        # Con el tipo y la traza: un StopIteration (o cualquier excepcion sin
        # mensaje) dejaba un "ERROR: " vacio que no decia nada (1-oct-2026).
        log(f"ERROR [{type(exc).__name__}]: {exc or 'sin mensaje'}")
        traceback.print_exc()
        sys.exit(1)
