"""
Calcula que deals ABIERTOS del equipo necesitan un GBS nuevo: los que salen
"SIN_LINK" en el ultimo chequeo de GBS Deal Checker (gbs-checker-YYYY-MM-DD.
xlsx, que ya solo trae deals abiertos) y que todavia no tienen ni un
borrador ni un GBS ya desplegado para ese mismo deal. Escribe el plan en
plan_creacion_gbs_deals.json. No toca vTiger ni SharePoint -- eso lo hace
Claude Code despues, siguiendo la seccion GENERAR del skill "gbs-deal-
maker" con este plan ya calculado (mismo patron "Calcular -> Ejecutar" que
Sofia/Calidad CRM/Notificar GBS).

Ejecutar con: python planificar_creacion_gbs_deals.py
"""
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

from paths import GBS_CHECKER_INFORMES_DIR as INFORMES_DIR, GBS_MAKER_DIR as MAKER_DIR

SCRIPT_DIR = Path(__file__).resolve().parent
BORRADORES_DIR = MAKER_DIR / "borradores"
DESPLEGADOS_DIR = MAKER_DIR / "desplegados"
PLAN_FILE = SCRIPT_DIR / "plan_creacion_gbs_deals.json"


def nombre_archivo(org, deal):
    base = re.sub(r'[\\/:*?"<>|]', "-", f"{org} - {deal}".strip())
    return base[:180] + ".docx"


def cargar_sin_link_hoy():
    hoy = datetime.now().strftime("%Y-%m-%d")
    archivo = INFORMES_DIR / f"gbs-checker-{hoy}.xlsx"
    if not archivo.exists():
        return None, archivo

    wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(header)}
    filas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue
        if row[idx["Estado GBS Link"]] != "SIN_LINK":
            continue
        filas.append({
            "owner": row[idx["Representante"]],
            "deal": row[idx["Deal"]],
            "org": row[idx["Organización"]],
            "etapa": row[idx["Etapa"]],
        })
    wb.close()
    return filas, archivo


def main():
    filas, archivo = cargar_sin_link_hoy()
    if filas is None:
        print(f"No hay un chequeo de GBS de hoy ({archivo.name} no existe). "
              "Corre 'Chequear GBS' primero.")
        sys.exit(1)

    ya_existe = set()
    for carpeta in (BORRADORES_DIR, DESPLEGADOS_DIR):
        if carpeta.exists():
            ya_existe.update(p.name for p in carpeta.iterdir() if p.suffix == ".docx")

    nuevos = []
    for f in filas:
        archivo_esperado = nombre_archivo(f["org"], f["deal"])
        if archivo_esperado in ya_existe:
            continue
        item = dict(f)
        item["archivo"] = archivo_esperado
        nuevos.append(item)

    with open(PLAN_FILE, "w", encoding="utf-8") as f:
        json.dump({"items": nuevos}, f, ensure_ascii=False, indent=2)

    print(f"Deals sin GBS Link hoy ({archivo.name}): {len(filas)}")
    print(f"Ya tienen borrador o GBS desplegado (se omiten): {len(filas) - len(nuevos)}")
    print(f"Nuevos por generar: {len(nuevos)}")
    for item in nuevos:
        print(f"  - {item['owner']}: {item['deal']} ({item['org']})")


if __name__ == "__main__":
    main()
